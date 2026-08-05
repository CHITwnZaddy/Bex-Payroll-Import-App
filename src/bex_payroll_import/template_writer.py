from __future__ import annotations

import re
import shutil
from copy import copy
from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.formula.translate import Translator

from bex_payroll_import.models import NormalizedPayrollRow, PU_SHEET_NAME, TDR_SHEET_NAME
from bex_payroll_import.spreadsheet_io import normalize_header


TDR_TEMPLATE_HEADERS = {"batch code", "eecode", "date", "earncode", "hours", "job"}
TDR_FORMULA_REFERENCE = re.compile(r"(?:'TDR'|TDR)!\$?[A-Z]+\$?(\d+)", re.IGNORECASE)


@dataclass(frozen=True)
class TemplateWriteResult:
    expected_output_rows: int


def copy_template_to_audit(template_path: Path, audit_path: Path) -> None:
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(template_path, audit_path)


def write_tdr_rows(audit_workbook_path: Path, rows: list[NormalizedPayrollRow]) -> TemplateWriteResult:
    workbook = load_workbook(audit_workbook_path)
    workbook.iso_dates = True
    try:
        if TDR_SHEET_NAME not in workbook.sheetnames:
            raise ValueError("Template is missing the TDR tab.")
        sheet = workbook[TDR_SHEET_NAME]
        data_start_row = find_tdr_data_start_row(sheet)
        formula_cells = capture_formula_cells(sheet, data_start_row)
        clear_existing_tdr_rows(sheet, data_start_row)

        for offset, payroll_row in enumerate(rows):
            target_row = data_start_row + offset
            write_tdr_source_cells(sheet, target_row, payroll_row)
            copy_formula_cells(sheet, formula_cells, target_row)

        result = inspect_pu_formula_coverage(workbook, data_start_row, len(rows))
        workbook.save(audit_workbook_path)
        return result
    finally:
        workbook.close()


def find_tdr_data_start_row(sheet) -> int:
    for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        headers = {normalize_header(value) for value in row if normalize_header(value)}
        if TDR_TEMPLATE_HEADERS.issubset(headers):
            return row_number + 1
    raise ValueError(
        "Template TDR tab has an unsupported layout. Expected Batch code, EECode, Date, EarnCode, Hours, and Job columns."
    )


def inspect_pu_formula_coverage(workbook, data_start_row: int, source_row_count: int) -> TemplateWriteResult:
    if PU_SHEET_NAME not in workbook.sheetnames:
        raise ValueError("Template is missing the PU tab.")

    source_rows = set(range(data_start_row, data_start_row + source_row_count))
    referenced_source_rows: list[int] = []
    pu_row_by_tdr_row: dict[int, list[int]] = {}
    pu_sheet = workbook[PU_SHEET_NAME]
    for row_number in range(1, pu_sheet.max_row + 1):
        formula = pu_sheet.cell(row=row_number, column=2).value
        if not isinstance(formula, str) or not formula.startswith("="):
            continue
        match = TDR_FORMULA_REFERENCE.search(formula)
        if match is None:
            continue
        tdr_row = int(match.group(1))
        if tdr_row in source_rows:
            referenced_source_rows.append(tdr_row)
            pu_row_by_tdr_row.setdefault(tdr_row, []).append(row_number)

    missing_rows = sorted(source_rows - set(referenced_source_rows))
    if missing_rows:
        if not referenced_source_rows:
            raise ValueError("Template PU tab has no formulas linked to the TDR data rows.")
        last_referenced_tdr_row = max(referenced_source_rows)
        expected_trailing_rows = list(range(last_referenced_tdr_row + 1, max(source_rows) + 1))
        if missing_rows != expected_trailing_rows:
            row_values = ", ".join(str(row) for row in missing_rows)
            raise ValueError(
                f"Template PU formulas skip TDR rows {row_values}. "
                "Correct the template formula sequence and rerun."
            )
        origin_pu_row = max(pu_row_by_tdr_row[last_referenced_tdr_row])
        extend_pu_formula_rows(
            pu_sheet,
            origin_pu_row=origin_pu_row,
            origin_tdr_row=last_referenced_tdr_row,
            missing_tdr_rows=missing_rows,
        )
        referenced_source_rows.extend(missing_rows)

    return TemplateWriteResult(expected_output_rows=len(referenced_source_rows))


def extend_pu_formula_rows(
    pu_sheet,
    *,
    origin_pu_row: int,
    origin_tdr_row: int,
    missing_tdr_rows: list[int],
) -> None:
    formula_cells = []
    for column in range(1, pu_sheet.max_column + 1):
        cell = pu_sheet.cell(row=origin_pu_row, column=column)
        if isinstance(cell.value, str) and cell.value.startswith("="):
            formula_cells.append((column, cell.value, cell.coordinate, copy(cell._style)))
    if not formula_cells:
        raise ValueError("Template PU tab has no formulas available to extend.")

    for tdr_row in missing_tdr_rows:
        target_pu_row = origin_pu_row + (tdr_row - origin_tdr_row)
        for column, formula, origin, style in formula_cells:
            target_cell = pu_sheet.cell(row=target_pu_row, column=column)
            target_cell.value = Translator(formula, origin=origin).translate_formula(target_cell.coordinate)
            target_cell._style = copy(style)


def clear_existing_tdr_rows(sheet, data_start_row: int) -> None:
    clear_columns = range(1, sheet.max_column + 1)
    for row in range(data_start_row, sheet.max_row + 1):
        for column in clear_columns:
            sheet.cell(row=row, column=column).value = None


def capture_formula_cells(sheet, data_start_row: int) -> dict[int, tuple[str, str]]:
    for row in range(data_start_row, sheet.max_row + 1):
        formula_cells = {}
        for column in range(1, sheet.max_column + 1):
            cell = sheet.cell(row=row, column=column)
            value = cell.value
            if isinstance(value, str) and value.startswith("="):
                formula_cells[column] = (value, cell.coordinate)
        if formula_cells:
            return formula_cells
    return {}


def write_tdr_source_cells(sheet, row_number: int, payroll_row: NormalizedPayrollRow) -> None:
    sheet.cell(row=row_number, column=1).value = payroll_row.batch_code
    sheet.cell(row=row_number, column=2).value = excel_lookup_code(payroll_row.employee_code)
    sheet.cell(row=row_number, column=8).value = payroll_row.work_date
    sheet.cell(row=row_number, column=12).value = payroll_row.department
    sheet.cell(row=row_number, column=13).value = excel_lookup_code(payroll_row.pay_type)
    sheet.cell(row=row_number, column=14).value = float(payroll_row.hours)
    sheet.cell(row=row_number, column=15).value = float(payroll_row.dollars or 0)
    sheet.cell(row=row_number, column=21).value = payroll_row.job
    sheet.cell(row=row_number, column=23).value = payroll_row.phase
    sheet.cell(row=row_number, column=25).value = payroll_row.cost_type


def copy_formula_cells(
    sheet,
    formula_cells: dict[int, tuple[str, str]],
    target_row: int,
) -> None:
    for column, (formula, origin) in formula_cells.items():
        target_cell = sheet.cell(row=target_row, column=column)
        target_cell.value = Translator(
            extend_pay_type_lookup_range(formula),
            origin=origin,
        ).translate_formula(target_cell.coordinate)


def excel_lookup_code(value: str) -> str | int:
    text = value.strip()
    if text.isdigit():
        return int(text)
    return text


def extend_pay_type_lookup_range(formula: str) -> str:
    return formula.replace("Key!$M$25:$N$33", "Key!$M$20:$N$33")
