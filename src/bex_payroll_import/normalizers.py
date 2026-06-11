from __future__ import annotations

from datetime import date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import load_workbook

from bex_payroll_import.key_lookup import EmployeeLookup, parse_expense_employee
from bex_payroll_import.models import NormalizedPayrollRow
from bex_payroll_import.spreadsheet_io import find_header_row, normalize_header


TDR_REQUIRED_COLUMNS = ["EECode", "Lastname", "Firstname", "Department", "EarnCode", "EarnHours"]
TDR_DATE_COLUMNS = ["Date", "InPunchTime", "OutPunchTime"]
EXPENSE_REQUIRED_COLUMNS = ["Employee", "Check Date", "Paid Amount"]


def normalize_tdr_file(tdr_path: Path, batch_code: str) -> list[NormalizedPayrollRow]:
    header_row = find_tdr_header_row(tdr_path)
    workbook = load_workbook(tdr_path, read_only=True, data_only=True)
    rows: list[NormalizedPayrollRow] = []
    try:
        sheet = workbook[header_row.sheet_name]
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row.row_number + 1, values_only=True),
            start=header_row.row_number + 1,
        ):
            employee_code = text_value(row, header_row.header_map, "EECode")
            if not employee_code:
                continue
            rows.append(
                NormalizedPayrollRow(
                    batch_code=batch_code,
                    employee_code=employee_code,
                    department=text_value(row, header_row.header_map, "Department"),
                    pay_type=text_value(row, header_row.header_map, "EarnCode"),
                    hours=decimal_value(cell(row, header_row.header_map, "EarnHours")),
                    job=text_value(row, header_row.header_map, "Job"),
                    phase=text_value(row, header_row.header_map, "Phase"),
                    cost_type=text_value(row, header_row.header_map, "Cost Type"),
                    work_date=date_value(first_available_cell(row, header_row.header_map, TDR_DATE_COLUMNS)),
                    dollars=decimal_or_none(cell(row, header_row.header_map, "Dollars")),
                    source="Time Detail Report",
                    source_row_number=row_number,
                )
            )
    finally:
        workbook.close()
    return rows


def load_employee_lookup_from_tdr(tdr_path: Path) -> EmployeeLookup:
    header_row = find_tdr_header_row(tdr_path)
    workbook = load_workbook(tdr_path, read_only=True, data_only=True)
    code_by_name: dict[tuple[str, str], str] = {}
    try:
        sheet = workbook[header_row.sheet_name]
        for row in sheet.iter_rows(min_row=header_row.row_number + 1, values_only=True):
            code = text_value(row, header_row.header_map, "EECode")
            first_name = text_value(row, header_row.header_map, "Firstname")
            last_name = text_value(row, header_row.header_map, "Lastname")
            if code and first_name and last_name:
                key = (normalize_lookup_name(first_name), normalize_lookup_name(last_name))
                existing_code = code_by_name.get(key)
                if existing_code and existing_code != code:
                    raise ValueError(
                        f"Duplicate employee name in Time Detail Report for {first_name} {last_name}: "
                        f"{existing_code} and {code}"
                    )
                code_by_name[key] = code
    finally:
        workbook.close()
    return EmployeeLookup(code_by_name)


def normalize_expense_file(
    expense_path: Path,
    lookup: EmployeeLookup,
    batch_code: str,
    default_check_date: date,
) -> list[NormalizedPayrollRow]:
    header_row = find_header_row(expense_path, EXPENSE_REQUIRED_COLUMNS)
    workbook = load_workbook(expense_path, read_only=True, data_only=True)
    rows: list[NormalizedPayrollRow] = []
    try:
        sheet = workbook[header_row.sheet_name]
        for row_number, row in enumerate(
            sheet.iter_rows(min_row=header_row.row_number + 1, values_only=True),
            start=header_row.row_number + 1,
        ):
            employee = text_value(row, header_row.header_map, "Employee")
            if not employee:
                continue
            identity = parse_expense_employee(employee)
            rows.append(
                NormalizedPayrollRow(
                    batch_code=batch_code,
                    employee_code=lookup.resolve_code(identity.first_name, identity.last_name),
                    department="DLPTO",
                    pay_type="EXP REIM",
                    hours=Decimal("0"),
                    job="",
                    phase="",
                    cost_type="L",
                    work_date=date_value_or_default(
                        cell(row, header_row.header_map, "Check Date"),
                        default_check_date,
                    ),
                    dollars=decimal_value(cell(row, header_row.header_map, "Paid Amount")),
                    source="Expense Transaction",
                    source_row_number=row_number,
                )
            )
    finally:
        workbook.close()
    return rows


def find_tdr_header_row(tdr_path: Path):
    header_row = find_header_row(tdr_path, TDR_REQUIRED_COLUMNS)
    if not any(normalize_header(column) in header_row.header_map for column in TDR_DATE_COLUMNS):
        raise ValueError("Missing required columns: Date, InPunchTime, or OutPunchTime")
    return header_row


def cell(row: tuple[object, ...], header_map: dict[str, int], column_name: str) -> object:
    index = header_map.get(normalize_header(column_name))
    if index is None or index > len(row):
        return None
    return row[index - 1]


def first_available_cell(row: tuple[object, ...], header_map: dict[str, int], column_names: list[str]) -> object:
    for column_name in column_names:
        value = cell(row, header_map, column_name)
        if value is not None and str(value).strip():
            return value
    return None


def text_value(row: tuple[object, ...], header_map: dict[str, int], column_name: str) -> str:
    value = cell(row, header_map, column_name)
    return "" if value is None else str(value).strip()


def decimal_value(value: object) -> Decimal:
    if value is None:
        return Decimal("0")
    text = str(value).strip()
    if not text:
        return Decimal("0")
    return Decimal(text).quantize(Decimal("0.01")).normalize()


def decimal_or_none(value: object) -> Decimal | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return decimal_value(text)


def date_value(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, str):
        text = value.strip()
        for date_format in ("%m/%d/%Y", "%Y-%m-%d %I:%M %p"):
            try:
                return datetime.strptime(text, date_format).date()
            except ValueError:
                pass
    raise ValueError(f"Expected date value, got {value!r}")


def date_value_or_default(value: object, default: date) -> date:
    if value is None:
        return default
    if isinstance(value, str) and not value.strip():
        return default
    return date_value(value)


def normalize_lookup_name(value: str) -> str:
    return " ".join(value.strip().lower().split())
