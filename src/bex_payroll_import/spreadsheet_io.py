from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path

from openpyxl import load_workbook


@dataclass(frozen=True)
class HeaderRow:
    sheet_name: str
    row_number: int
    header_map: dict[str, int]


def normalize_header(value: object) -> str:
    return " ".join(str(value or "").strip().lower().split())


def build_header_map(headers: list[object]) -> dict[str, int]:
    return {
        normalize_header(header): index
        for index, header in enumerate(headers, start=1)
        if normalize_header(header)
    }


def find_header_row(workbook_path: Path, required_columns: list[str]) -> HeaderRow:
    workbook = load_workbook(workbook_path, read_only=True, data_only=False)
    normalized_required = [normalize_header(column) for column in required_columns]
    try:
        for sheet in workbook.worksheets:
            for row_number, row in enumerate(sheet.iter_rows(values_only=True), start=1):
                header_map = build_header_map(list(row))
                if all(column in header_map for column in normalized_required):
                    return HeaderRow(sheet.title, row_number, header_map)
        missing = ", ".join(required_columns)
        raise ValueError(f"Missing required columns: {missing}")
    finally:
        workbook.close()
