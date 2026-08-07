from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from bex_payroll_import.models import EmployeeIdentity, KEY_SHEET_NAME
from bex_payroll_import.spreadsheet_io import build_header_map, normalize_header


class EmployeeLookup:
    def __init__(
        self,
        code_by_name: dict[tuple[str, str], str],
        fallback: EmployeeLookup | None = None,
    ) -> None:
        self._code_by_name = code_by_name
        self._fallback = fallback

    def resolve_code(self, first_name: str, last_name: str) -> str:
        key = (normalize_name(first_name), normalize_name(last_name))
        if key in self._code_by_name:
            return self._code_by_name[key]
        if self._fallback is not None:
            return self._fallback.resolve_code(first_name, last_name)
        raise KeyError(f"No employee code found for {first_name} {last_name}")

    def with_fallback(self, fallback: EmployeeLookup) -> EmployeeLookup:
        return EmployeeLookup(self._code_by_name, fallback=fallback)


def normalize_name(value: str) -> str:
    return " ".join(value.strip().lower().split())


def display_name(value: str) -> str:
    return " ".join(piece.capitalize() for piece in value.strip().split())


def parse_expense_employee(raw_employee: object) -> EmployeeIdentity:
    raw = str(raw_employee or "").strip()
    if "," not in raw:
        raise ValueError(f"Expected employee format 'LAST, FIRST', got '{raw}'")
    last_name, first_name = [part.strip() for part in raw.split(",", 1)]
    if not first_name or not last_name:
        raise ValueError(f"Expected employee format 'LAST, FIRST', got '{raw}'")
    return EmployeeIdentity(first_name=display_name(first_name), last_name=display_name(last_name))


def load_employee_lookup(template_path: Path) -> EmployeeLookup:
    workbook = load_workbook(template_path, read_only=True, data_only=False)
    try:
        if KEY_SHEET_NAME not in workbook.sheetnames:
            raise ValueError("Template is missing the Key tab.")
        sheet = workbook[KEY_SHEET_NAME]
        header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
        header_map = build_header_map(list(header_row))
        code_col = find_first_header(header_map, ["employee code", "eecode", "ee code"])
        last_col = find_first_header(header_map, ["last name", "lastname"])
        first_col = find_first_header(header_map, ["first name", "firstname"])

        code_by_name: dict[tuple[str, str], str] = {}
        for row in sheet.iter_rows(min_row=2, values_only=True):
            code = cell_value(row, code_col)
            first_name = cell_value(row, first_col)
            last_name = cell_value(row, last_col)
            code_text = str(code).strip() if code is not None else ""
            first_key = normalize_name(str(first_name)) if first_name is not None else ""
            last_key = normalize_name(str(last_name)) if last_name is not None else ""
            if code_text and first_key and last_key:
                key = (first_key, last_key)
                existing_code = code_by_name.get(key)
                if existing_code and existing_code != code_text:
                    raise ValueError(
                        f"Duplicate employee name in Key tab for {first_name} {last_name}: "
                        f"{existing_code} and {code_text}"
                    )
                code_by_name[key] = code_text
        return EmployeeLookup(code_by_name)
    finally:
        workbook.close()


def find_first_header(header_map: dict[str, int], options: list[str]) -> int:
    for option in options:
        normalized = normalize_header(option)
        if normalized in header_map:
            return header_map[normalized]
    raise ValueError(f"Missing Key tab column. Tried: {', '.join(options)}")


def cell_value(row: tuple[object, ...], one_based_index: int) -> object:
    return row[one_based_index - 1] if 1 <= one_based_index <= len(row) else None
