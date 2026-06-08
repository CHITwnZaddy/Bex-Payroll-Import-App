from datetime import date, datetime
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

import bex_payroll_import.runner as runner
from bex_payroll_import.models import RunInputs
from bex_payroll_import.runner import run_payroll_import
from tests.helpers import save_workbook


class FakeExcelBackend:
    def recalculate_and_export(self, workbook_path: Path, csv_path: Path) -> None:
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        csv_path.write_text("CD0529143708,E100,DLPTO,REG,8,100,200,L,04/12/2026,0\n", encoding="utf-8")


class HeaderOnlyExcelBackend:
    def recalculate_and_export(self, workbook_path: Path, csv_path: Path) -> None:
        csv_path.parent.mkdir(parents=True, exist_ok=True)
        csv_path.write_text("Batch code,Employee code\n", encoding="utf-8")


def make_template(path: Path) -> Path:
    workbook = Workbook()
    pu = workbook.active
    pu.title = "PU"
    workbook.create_sheet("Key")
    tdr = workbook.create_sheet("TDR")
    tdr.append(["Batch code", "EECode", "Lastname", "Firstname", "HomeDepartment", "Pay Class", "Badge", "Date", "Text to Col", "InPunchTime", "OutPunchTime", "Department", "EarnCode", "Hours", "Dollars", "Employee Approved", "Supervisor Approved", "Tax Profile", "Home Department Desc", "Dist Department Desc", "Job", "Employee code", "Phase", "Cost Code", "Cost type", "Home Department", "Department", "Phase Adj", "Phase", "Pay type"])
    tdr.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "=B2", "=W2", "", "=Y2", "", "=L2", "", "=W2", "=M2"])
    key = workbook["Key"]
    key.append(["Employee code", "Last Name", "First Name"])
    key.append(["E100", "Doe", "Jane"])
    workbook.save(path)
    return path


def make_template_with_real_key_shape(path: Path) -> Path:
    workbook = Workbook()
    pu = workbook.active
    pu.title = "PU"
    key = workbook.create_sheet("Key")
    key.append(["Code", "Dep", "Office", None, "Foreman", "Phase", "Category", None, "Code", "Name"])
    key.append(["BAKKEN", "6", "N", None, "BAKKEN", 1, 1010, None, 48, "BAKKEN"])
    tdr = workbook.create_sheet("TDR")
    tdr.append(["Batch code", "EECode", "Lastname", "Firstname", "HomeDepartment", "Pay Class", "Badge", "Date", "Text to Col", "InPunchTime", "OutPunchTime", "Department", "EarnCode", "Hours", "Dollars", "Employee Approved", "Supervisor Approved", "Tax Profile", "Home Department Desc", "Dist Department Desc", "Job", "Employee code", "Phase", "Cost Code", "Cost type", "Home Department", "Department", "Phase Adj", "Phase", "Pay type"])
    tdr.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "=B2", "=W2", "", "=Y2", "", "=L2", "", "=W2", "=M2"])
    workbook.save(path)
    return path


def make_template_without_tdr(path: Path) -> Path:
    workbook = Workbook()
    pu = workbook.active
    pu.title = "PU"
    key = workbook.create_sheet("Key")
    key.append(["Employee code", "Last Name", "First Name"])
    key.append(["E100", "Doe", "Jane"])
    workbook.save(path)
    return path


def cell_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raise AssertionError(f"Expected date-like value, got {value!r}")


def test_run_payroll_import_success_with_expenses(tmp_path: Path) -> None:
    template_path = make_template(tmp_path / "template.xlsx")
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    tdr_path = save_workbook(
        source_dir / "tdr.xlsx",
        "Time Detail Repor",
        [["EECode", "Lastname", "Firstname", "Date", "Department", "EarnCode", "EarnHours"], ["E100", "Doe", "Jane", "04/12/2026", "DLPTO", "REG", 8]],
    )
    expense_path = save_workbook(
        source_dir / "expense.xlsx",
        "Expense Transacti",
        [["Employee", "Expense Date", "Paid Amount"], ["DOE, JANE", "04/15/2026", 55]],
    )

    outputs = run_payroll_import(
        RunInputs(template_path=template_path, tdr_path=tdr_path, expense_path=expense_path, expense_skipped=False),
        now=datetime(2026, 5, 29, 14, 37, 8),
        excel_backend=FakeExcelBackend(),
    )

    assert outputs.batch_code == "CD0529143708"
    assert outputs.payroll_csv_path is not None
    assert outputs.payroll_csv_path.exists()
    assert outputs.audit_workbook_path.exists()
    assert outputs.validation_path.exists()

    workbook = load_workbook(outputs.audit_workbook_path, data_only=False)
    try:
        tdr = workbook["TDR"]
        assert tdr["B3"].value == "E100"
        assert cell_date(tdr["H3"].value) == date(2026, 4, 15)
        assert tdr["L3"].value == "DLPTO"
        assert tdr["M3"].value == "EXP REIM"
        assert tdr["N3"].value == 0
        assert tdr["O3"].value == 55
        assert tdr["Y3"].value == "L"
    finally:
        workbook.close()


def test_run_payroll_import_uses_tdr_employee_codes_for_expenses(tmp_path: Path) -> None:
    template_path = make_template_with_real_key_shape(tmp_path / "template.xlsx")
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    tdr_path = save_workbook(
        source_dir / "tdr.xlsx",
        "Time Detail Repor",
        [
            [
                "EECode",
                "Lastname",
                "Firstname",
                "HomeDepartment",
                "Pay Class",
                "Badge",
                "InPunchTime",
                "OutPunchTime",
                "Department",
                "EarnCode",
                "EarnHours",
                "Dollars",
            ],
            ["0048", "BAKER", "KENNETH", "200", "SAL", "0049", "2026-04-13 12:00 AM", "2026-04-13 12:00 AM", "9 Mile improvements", "R", 8, 0],
            ["1009", "LUMLEY", "MICHAEL", "200", "SAL", "1010", "2026-04-13 12:00 AM", "2026-04-13 12:00 AM", "9 Mile improvements", "R", 8, 0],
        ],
    )
    expense_path = save_workbook(
        source_dir / "expense.xlsx",
        "Expense Transacti",
        [
            ["Employee", "Expense Date", "Paid Amount"],
            ["LUMLEY, MICHAEL", "04/13/2026", 97.14],
            ["BAKER, KENNETH", "04/17/2026", 41.61],
        ],
    )

    outputs = run_payroll_import(
        RunInputs(template_path=template_path, tdr_path=tdr_path, expense_path=expense_path, expense_skipped=False),
        now=datetime(2026, 5, 29, 14, 37, 8),
        excel_backend=FakeExcelBackend(),
    )

    assert outputs.payroll_csv_path is not None
    workbook = load_workbook(outputs.audit_workbook_path, data_only=False)
    try:
        tdr = workbook["TDR"]
        assert tdr["B4"].value == "1009"
        assert tdr["M4"].value == "EXP REIM"
        assert tdr["O4"].value == 97.14
        assert tdr["B5"].value == "0048"
        assert tdr["M5"].value == "EXP REIM"
        assert tdr["O5"].value == 41.61
    finally:
        workbook.close()


def test_run_payroll_import_hard_stops_when_expense_missing_and_not_confirmed(tmp_path: Path) -> None:
    template_path = make_template(tmp_path / "template.xlsx")
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    tdr_path = save_workbook(source_dir / "tdr.xlsx", "Time Detail Repor", [["EECode", "Lastname", "Firstname", "Date", "Department", "EarnCode", "EarnHours"]])

    outputs = run_payroll_import(
        RunInputs(template_path=template_path, tdr_path=tdr_path, expense_path=None, expense_skipped=False),
        now=datetime(2026, 5, 29, 14, 37, 8),
        excel_backend=FakeExcelBackend(),
    )

    assert outputs.payroll_csv_path is None
    assert outputs.validation_path.name.startswith("ValidationErrors_")
    assert outputs.validation_path.exists()


def test_run_payroll_import_reports_template_write_failure(tmp_path: Path) -> None:
    template_path = make_template_without_tdr(tmp_path / "template.xlsx")
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    tdr_path = save_workbook(
        source_dir / "tdr.xlsx",
        "Time Detail Repor",
        [["EECode", "Lastname", "Firstname", "Date", "Department", "EarnCode", "EarnHours"], ["E100", "Doe", "Jane", "04/12/2026", "DLPTO", "REG", 8]],
    )

    outputs = run_payroll_import(
        RunInputs(template_path=template_path, tdr_path=tdr_path, expense_path=None, expense_skipped=True),
        now=datetime(2026, 5, 29, 14, 37, 8),
        excel_backend=FakeExcelBackend(),
    )

    assert outputs.payroll_csv_path is None
    assert outputs.validation_path.name.startswith("ValidationErrors_")
    assert outputs.validation_path.exists()
    assert "Template is missing the TDR tab." in outputs.validation_path.read_text(encoding="utf-8")


def test_run_payroll_import_reports_template_copy_failure(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    template_path = make_template(tmp_path / "template.xlsx")
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    tdr_path = save_workbook(
        source_dir / "tdr.xlsx",
        "Time Detail Repor",
        [["EECode", "Lastname", "Firstname", "Date", "Department", "EarnCode", "EarnHours"], ["E100", "Doe", "Jane", "04/12/2026", "DLPTO", "REG", 8]],
    )

    def fail_copy(template_path: Path, audit_path: Path) -> None:
        raise PermissionError("template copy failed")

    monkeypatch.setattr(runner, "copy_template_to_audit", fail_copy)

    outputs = run_payroll_import(
        RunInputs(template_path=template_path, tdr_path=tdr_path, expense_path=None, expense_skipped=True),
        now=datetime(2026, 5, 29, 14, 37, 8),
        excel_backend=FakeExcelBackend(),
    )

    assert outputs.payroll_csv_path is None
    assert outputs.validation_path.name.startswith("ValidationErrors_")
    assert outputs.validation_path.exists()
    assert "template copy failed" in outputs.validation_path.read_text(encoding="utf-8")


def test_run_payroll_import_preserves_exported_csv_when_final_validation_fails(tmp_path: Path) -> None:
    template_path = make_template(tmp_path / "template.xlsx")
    source_dir = tmp_path / "source"
    source_dir.mkdir()
    tdr_path = save_workbook(
        source_dir / "tdr.xlsx",
        "Time Detail Repor",
        [["EECode", "Lastname", "Firstname", "Date", "Department", "EarnCode", "EarnHours"], ["E100", "Doe", "Jane", "04/12/2026", "DLPTO", "REG", 8]],
    )

    outputs = run_payroll_import(
        RunInputs(template_path=template_path, tdr_path=tdr_path, expense_path=None, expense_skipped=True),
        now=datetime(2026, 5, 29, 14, 37, 8),
        excel_backend=HeaderOnlyExcelBackend(),
    )

    assert outputs.payroll_csv_path is not None
    assert outputs.payroll_csv_path.exists()
    assert outputs.validation_path.name.startswith("ValidationErrors_")
    assert outputs.validation_path.exists()
