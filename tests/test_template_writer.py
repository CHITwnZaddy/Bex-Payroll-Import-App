from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook, load_workbook

from bex_payroll_import.models import NormalizedPayrollRow
from bex_payroll_import.template_writer import copy_template_to_audit, write_tdr_rows


def make_payroll_row(
    employee_code: str = "E100",
    job: str = "100",
    phase: str = "200",
    cost_type: str = "L",
) -> NormalizedPayrollRow:
    return NormalizedPayrollRow(
        "CD0529143708",
        employee_code,
        "DLPTO",
        "REG",
        Decimal("8"),
        job,
        phase,
        cost_type,
        date(2026, 4, 12),
        Decimal("0"),
        "Time Detail Report",
        2,
    )


def make_template(path: Path) -> Path:
    workbook = Workbook()
    pu = workbook.active
    pu.title = "PU"
    pu.append(["Batch code", "Employee code"])
    for row_number in range(2, 12):
        pu.cell(row=row_number, column=2).value = f"=TDR!V{row_number}"
    tdr = workbook.create_sheet("TDR")
    workbook.create_sheet("Key")
    tdr.append(["Batch code", "EECode", "Lastname", "Firstname", "HomeDepartment", "Pay Class", "Badge", "Date", "Text to Col", "InPunchTime", "OutPunchTime", "Department", "EarnCode", "Hours", "Dollars", "Employee Approved", "Supervisor Approved", "Tax Profile", "Home Department Desc", "Dist Department Desc", "Job", "Employee code", "Phase", "Cost Code", "Cost type", "Home Department", "Department", "Phase Adj", "Phase", "Pay type"])
    tdr.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "=B2", "=W2", "", "=Y2", "", "=L2", "", "=W2", '=IFERROR(VLOOKUP(M2,Key!$M$25:$N$33,2,FALSE),"NL")'])
    workbook.save(path)
    return path


def make_real_template_shape(path: Path) -> Path:
    workbook = Workbook()
    pu = workbook.active
    pu.title = "PU"
    pu.append(["Batch code", "Employee code"])
    pu.cell(row=2, column=2).value = "=TDR!V3"
    pu.cell(row=3, column=2).value = "=TDR!V3"
    pu.cell(row=4, column=2).value = "=TDR!V4"
    pu.cell(row=4, column=9).value = "=TDR!H4"
    pu.cell(row=4, column=9).number_format = "m/d/yyyy"
    pu.cell(row=4, column=22).value = "=TDR!O4"
    pu.cell(row=4, column=22).number_format = "0.00"
    tdr = workbook.create_sheet("TDR")
    workbook.create_sheet("Key")
    tdr.append([None, None, None, None, None, None, None, None, "ADD"])
    tdr.append(["Batch code", "EECode", "Lastname", "Firstname", "HomeDepartment", "Pay Class", "Badge", "Date", "Text to Col", "InPunchTime", "OutPunchTime", "Department", "EarnCode", "Hours", "Dollars", "Employee Approved", "Supervisor Approved", "Tax Profile", "Home Department Desc", "Dist Department Desc", "Job", "Employee code", "Phase", "Cost Code", "Cost type", "Home Department", "Department", "Phase Adj", "Phase", "Pay type"])
    tdr.append(["", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "", "=B3", "=W3", "", "=Y3", "", "=L3", "", "=W3", '=IFERROR(VLOOKUP(M3,Key!$M$25:$N$33,2,FALSE),"NL")'])
    workbook.save(path)
    return path


def test_copy_template_to_audit_preserves_original(tmp_path: Path) -> None:
    template_path = make_template(tmp_path / "template.xlsx")
    audit_path = tmp_path / "out" / "audit.xlsx"

    copy_template_to_audit(template_path, audit_path)
    write_tdr_rows(audit_path, [make_payroll_row()])

    assert audit_path.exists()
    original_workbook = load_workbook(template_path, data_only=False)
    audit_workbook = load_workbook(audit_path, data_only=False)
    assert original_workbook["TDR"]["A2"].value is None
    assert original_workbook["TDR"]["V2"].value == "=B2"
    assert audit_workbook["TDR"]["A2"].value == "CD0529143708"


def test_write_tdr_rows_adds_rows_and_copies_formula_columns(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_template(audit_path)
    rows = [
        NormalizedPayrollRow("CD0529143708", "E100", "DLPTO", "REG", Decimal("8"), "100", "200", "L", date(2026, 4, 12), Decimal("0"), "Time Detail Report", 2),
        NormalizedPayrollRow("CD0529143708", "E101", "DLPTO", "EXP REIM", Decimal("0"), "", "", "L", date(2026, 4, 15), Decimal("55.25"), "Expense Transaction", 2),
    ]

    write_tdr_rows(audit_path, rows)

    workbook = load_workbook(audit_path, data_only=False)
    sheet = workbook["TDR"]
    assert sheet["A2"].value == "CD0529143708"
    assert sheet["B2"].value == "E100"
    assert sheet["H2"].value == date(2026, 4, 12)
    assert sheet["N3"].value == 0
    assert sheet["O3"].value == 55.25
    assert sheet["V3"].value == "=B3"


def test_write_tdr_rows_preserves_real_template_headers_and_starts_on_row_three(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_real_template_shape(audit_path)

    result = write_tdr_rows(audit_path, [make_payroll_row(), make_payroll_row("E101")])

    workbook = load_workbook(audit_path, data_only=False)
    try:
        sheet = workbook["TDR"]
        assert sheet["A2"].value == "Batch code"
        assert sheet["B2"].value == "EECode"
        assert sheet["A3"].value == "CD0529143708"
        assert sheet["B3"].value == "E100"
        assert sheet["V3"].value == "=B3"
        assert sheet["A4"].value == "CD0529143708"
        assert sheet["B4"].value == "E101"
        assert sheet["V4"].value == "=B4"
        assert result.expected_output_rows == 3
    finally:
        workbook.close()


def test_write_tdr_rows_extends_pu_formulas_to_cover_additional_source_rows(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_real_template_shape(audit_path)

    result = write_tdr_rows(
        audit_path,
        [make_payroll_row(), make_payroll_row("E101"), make_payroll_row("E102")],
    )

    workbook = load_workbook(audit_path, data_only=False)
    try:
        pu = workbook["PU"]
        assert pu["B5"].value == "=TDR!V5"
        assert pu["I5"].value == "=TDR!H5"
        assert pu["I5"].number_format == "m/d/yyyy"
        assert pu["V5"].value == "=TDR!O5"
        assert pu["V5"].number_format == "0.00"
        assert result.expected_output_rows == 4
    finally:
        workbook.close()


def test_write_tdr_rows_writes_numeric_employee_codes_for_excel_lookup(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_template(audit_path)

    write_tdr_rows(audit_path, [make_payroll_row("0048")])

    workbook = load_workbook(audit_path, data_only=False)
    sheet = workbook["TDR"]
    assert sheet["B2"].value == 48


def test_write_tdr_rows_writes_numeric_earn_code_for_excel_lookup(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_template(audit_path)
    payroll_row = make_payroll_row()
    payroll_row = NormalizedPayrollRow(
        payroll_row.batch_code,
        payroll_row.employee_code,
        "DLPTO",
        "5",
        payroll_row.hours,
        payroll_row.job,
        payroll_row.phase,
        payroll_row.cost_type,
        payroll_row.work_date,
        payroll_row.dollars,
        payroll_row.source,
        payroll_row.source_row_number,
    )

    write_tdr_rows(audit_path, [payroll_row])

    workbook = load_workbook(audit_path, data_only=False)
    sheet = workbook["TDR"]
    assert sheet["M2"].value == 5


def test_write_tdr_rows_clears_stale_rows_below_current_import(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_template(audit_path)
    workbook = load_workbook(audit_path)
    sheet = workbook["TDR"]
    sheet.append(["OLD", "OLD", "OLD", "OLD", "OLD", "OLD", "OLD", "OLD", "OLD", "OLD", "OLD", "OLD", "OLD", 1, 2, "", "", "", "", "", "OLD", "=B3", "OLD", "", "OLD", "", "=L3", "", "=W3", "=M3"])
    sheet.append(["OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", "OLDER", 3, 4, "", "", "", "", "", "OLDER", "=B4", "OLDER", "", "OLDER", "", "=L4", "", "=W4", "=M4"])
    workbook.save(audit_path)

    write_tdr_rows(audit_path, [make_payroll_row()])

    workbook = load_workbook(audit_path, data_only=False)
    sheet = workbook["TDR"]
    stale_columns = [*range(1, 16), *range(21, 31)]
    assert [sheet.cell(row=3, column=column).value for column in stale_columns] == [None] * len(stale_columns)


def test_write_tdr_rows_preserves_job_and_copies_phase_formula_columns(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_template(audit_path)
    rows = [
        make_payroll_row("E100", "100", "200", "L"),
        make_payroll_row("E101", "300", "400", "M"),
    ]

    write_tdr_rows(audit_path, rows)

    workbook = load_workbook(audit_path, data_only=False)
    sheet = workbook["TDR"]
    assert sheet["U2"].value == "100"
    assert sheet["W2"].value == "=W2"
    assert sheet["Y2"].value == "=Y2"
    assert sheet["U3"].value == "300"
    assert sheet["W3"].value == "=W3"
    assert sheet["Y3"].value == "=Y3"


def test_write_tdr_rows_translates_all_formula_columns(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_template(audit_path)

    write_tdr_rows(audit_path, [make_payroll_row(), make_payroll_row("E101")])

    workbook = load_workbook(audit_path, data_only=False)
    sheet = workbook["TDR"]
    assert sheet["AA3"].value == "=L3"
    assert sheet["AC3"].value == "=W3"
    assert sheet["AD3"].value == '=IFERROR(VLOOKUP(M3,Key!$M$20:$N$33,2,FALSE),"NL")'


def test_write_tdr_rows_extends_pay_type_lookup_range_for_earn_code_5(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    make_template(audit_path)

    write_tdr_rows(audit_path, [make_payroll_row()])

    workbook = load_workbook(audit_path, data_only=False)
    sheet = workbook["TDR"]
    assert sheet["AD2"].value == '=IFERROR(VLOOKUP(M2,Key!$M$20:$N$33,2,FALSE),"NL")'


def test_write_tdr_rows_requires_tdr_sheet(tmp_path: Path) -> None:
    audit_path = tmp_path / "audit.xlsx"
    workbook = Workbook()
    workbook.active.title = "PU"
    workbook.save(audit_path)

    try:
        write_tdr_rows(audit_path, [make_payroll_row()])
    except ValueError as error:
        assert str(error) == "Template is missing the TDR tab."
    else:
        raise AssertionError("Expected missing TDR tab to raise ValueError.")
